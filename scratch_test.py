import os
import docx
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, nsmap

nsmap['v'] = 'urn:schemas-microsoft-com:vml'
nsmap['o'] = 'urn:schemas-microsoft-com:office:office'

def test():
    doc = Document()
    header = doc.sections[0].header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    watermark_path = os.path.abspath('d:/Coding/django learning/excelDjango1/shirva_report_system/reports/assets/watermark_fade.png')
    
    rId, _ = header.part.get_or_add_image(watermark_path)

    vml_xml = f'''<w:r {nsdecls("w")}>
        <w:rPr><w:noProof/></w:rPr>
        <w:pict {nsdecls("v", "o")}>
            <v:shape id="WaterMark" style="position:absolute;margin-left:0;margin-top:0;width:450pt;height:450pt;z-index:-251657728;mso-position-horizontal:center;mso-position-horizontal-relative:margin;mso-position-vertical:center;mso-position-vertical-relative:margin" coordsize="21600,21600" fillcolor="none" stroked="f">
                <v:imagedata r:id="{rId}" {nsdecls("r")}/>
            </v:shape>
        </w:pict>
    </w:r>'''
    hp._p.append(parse_xml(vml_xml))

    logo_path = os.path.abspath('d:/Coding/django learning/excelDjango1/shirva_report_system/reports/assets/karnataka_logo.png')
    r_logo = hp.add_run()
    r_logo.add_picture(logo_path, width=Inches(1.0))

    doc.save('test_out.docx')
    print('Test docx with watermark and header logo saved successfully!')

if __name__ == "__main__":
    test()
